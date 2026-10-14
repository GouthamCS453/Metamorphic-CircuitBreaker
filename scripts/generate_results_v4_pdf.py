import sys
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_honest_report():
    doc = Document()
    
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("Metamorphic Circuit Breaker (MCB) Framework")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 51, 102)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Critical Verification & Scientific Results Report (Protocol v4)\nUnvarnished Audit, Baseline Domination Analysis, and Negative Findings")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(110, 110, 110)
    
    # Metadata Box
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Evaluation Cohort", "ISIC Partition B (N=125, 21 errors)"), ("Audit Status", "Strict Protocol v4 Compliant"), ("Pre-Registered D1-D4", "0 Met, 3 Not Met, 1 Unverified")],
        [("Comparative Outcome", "TTA / k-of-n Dominates MCB"), ("Attribution Test", "COLLAPSED to Geometric (28.3%)"), ("Hardware Accounting", "CPU Execution (2,613 ms/image)")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (k, v) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r_k = p.add_run(f"{k}: ")
            r_k.font.bold = True
            r_k.font.size = Pt(8.5)
            r_v = p.add_run(v)
            r_v.font.size = Pt(8.5)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # ── Section 1: Executive Audit Summary ────────────────────────────────────
    h1 = doc.add_heading("1. Executive Summary & Critical Scientific Findings", level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    
    p = doc.add_paragraph()
    p.add_run("Following rigorous verification of Protocol v4 on held-out Partition B (N=125, 21 errors), this report documents the true experimental outcomes without spin. While the underlying confusion matrix arithmetic and Wilson intervals are mathematically sound, the empirical results decisively contradict several core marketing claims:")
    
    bullets = [
        ("TTA Variance and k-of-n Flip Fractions Dominate MCB: ", "At op_far10, TTA prediction variance achieves 52.4% Recall at 5.8% FAR (11 TP, 6 FP). MCB achieves the exact same 52.4% Recall but incurs 8.7% FAR (11 TP, 9 FP)—representing 3 additional false alarms with zero additional caught errors. At op_j, TTA catches 13 errors (61.9% recall) vs MCB's 11 errors (52.4% recall) at the exact same 8.7% FAR."),
        ("Pre-Registered Decision Rules Failed: ", "D1 is NOT MET (AUROC 95% CI spans zero [-0.019, +0.085], and TTA beats MCB at matched FAR). D2 is NOT MET (Hierarchical CBI ties Flat MT: 0.864 vs 0.865; hierarchy adds no discriminative value). D4 is NOT MET (attribution accuracy was 28.3%, indistinguishable from 33.3% random chance; p=0.25)."),
        ("Attribution Collapsed to Geometric Bias: ", "Across 60 synthetic injections, the argmax decision rule predicted the Geometric family 42 times (70.0%), regardless of whether severe photometric or sensor faults were injected. Semantic fault attribution was NOT demonstrated."),
        ("Prior Default is Weak: ", "The heuristic prior (theta=0.55) catches only 5 of 21 errors (23.8% recall, CI [10.6, 45.1]), missing over 76% of model failures. It is an uncalibrated, high-precision, low-recall operating point, not a reliable clinical safety interceptor.")
    ]
    for b_title, b_text in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(2)
        r_bt = bp.add_run(b_title)
        r_bt.font.bold = True
        bp.add_run(b_text)
        
    # ── Section 2: Rigorous Master Comparison Table ───────────────────────────
    h2 = doc.add_heading("2. Master Benchmark Comparison (Partition B Held-Out)", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    
    table_data = [
        ["Monitoring Paradigm", "Operating Point / Cutoff", "TP", "FP", "TN", "FN", "Safety Recall (95% CI)", "False Alarm Rate (95% CI)", "Youden J", "AUROC (95% CI)"],
        ["Standalone (Unmonitored)", "None (raw model)", "0", "0", "104", "21", "0.0%", "0.0%", "0.000", "—"],
        ["Naive 1-Flip MT", "Strawman (k >= 1)", "15", "20", "84", "6", "71.4% [50.0, 86.2]", "19.2% [12.8, 27.9]", "0.522", "—"],
        ["Continuous k-of-n Flips", "k >= 3 (of 25)", "11", "6", "98", "10", "52.4% [32.4, 71.7]", "5.8% [2.7, 12.0]", "0.466", "0.798 [0.681, 0.915]"],
        ["TTA Prediction Variance", "op_far10 (thr=0.115)", "11", "6", "98", "10", "52.4% [32.4, 71.7]", "5.8% [2.7, 12.0]", "0.466", "0.798 [0.681, 0.915]"],
        ["TTA Prediction Variance", "op_j (thr=0.077)", "13", "9", "95", "8", "61.9% [40.9, 79.3]", "8.7% [4.6, 15.6]", "0.532", "0.798 [0.681, 0.915]"],
        ["Softmax Entropy", "op_far10 (thr=1.385)", "10", "11", "93", "11", "47.6% [28.3, 67.6]", "10.6% [6.0, 18.0]", "0.370", "0.835 [0.750, 0.911]"],
        ["Softmax Entropy", "op_j (thr=1.102)", "17", "24", "80", "4", "81.0% [60.0, 92.3]", "23.1% [16.0, 32.1]", "0.579", "0.835 [0.750, 0.911]"],
        ["MC-Dropout (Head-Only)", "op_far10 (thr=1.377)", "11", "12", "92", "10", "52.4% [32.4, 71.7]", "11.5% [6.7, 19.1]", "0.408", "0.827 [0.739, 0.905]"],
        ["MC-Dropout (Head-Only)", "op_j (thr=1.111)", "15", "25", "79", "6", "71.4% [50.0, 86.2]", "24.0% [16.9, 33.1]", "0.474", "0.827 [0.739, 0.905]"],
        ["MCB Prior Default", "Heuristic (tau=0.35, theta=0.55)", "5", "1", "103", "16", "23.8% [10.6, 45.1]", "1.0% [0.2, 5.3]", "0.229", "0.860 [0.777, 0.938]"],
        ["MCB Calibrated op_far10", "Frozen (tau=0.10, theta=0.25)", "11", "9", "95", "10", "52.4% [32.4, 71.7]", "8.7% [4.6, 15.6]", "0.437", "0.864 [0.782, 0.942]"],
        ["MCB Calibrated op_j", "Frozen (tau=0.10, theta=0.25)", "11", "9", "95", "10", "52.4% [32.4, 71.7]", "8.7% [4.6, 15.6]", "0.437", "0.864 [0.782, 0.942]"]
    ]
    
    t = doc.add_table(rows=len(table_data), cols=10)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_vals in enumerate(table_data):
        for c_idx, val in enumerate(row_vals):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            elif "TTA" in row_vals[0] and "op_far10" in row_vals[1]:
                set_cell_background(cell, "EAF7EC") # Green highlight for winning baseline
                run.font.bold = True
            elif "MCB" in row_vals[0]:
                set_cell_background(cell, "F2F5F8")
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # ── Section 3: Honest Pre-Registered Decision Rules (D1 - D4) ─────────────
    h3 = doc.add_heading("3. Pre-Registered Decision Rules (Protocol v4 Verdict)", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    
    rules_table = [
        ["Rule ID", "Pre-Registered Hypothesis", "Evaluated Finding", "Audit Verdict", "Scientific Explanation"],
        ["D1", "MCB significantly beats best baseline on AUROC or matched FAR recall", "Paired Bootstrap AUC Diff: +0.029 [-0.019, +0.085]\nMatched FAR Recall: TTA achieves 52.4% @ 5.8% FAR vs MCB 52.4% @ 8.7% FAR", "NOT MET", "AUROC CI includes zero. TTA achieves identical recall with 3 fewer false alarms (5.8% vs 8.7% FAR). MCB is dominated by TTA."],
        ["D2", "Hierarchical CBI significantly beats Flat MT", "AUROC Diff: -0.0014 [-0.0076, +0.0044]\nFlat AUC: 0.865 vs Hierarchical: 0.864", "NOT MET (TIE)", "Flat MT and Hierarchical CBI produce virtually identical discrimination when confidence weighting is included. Hierarchy adds zero detection value."],
        ["D3", "MCB catches overconfident silent misclassifications (c0 >= 0.90)", "Errors with c0 >= 0.90: 0\nTop-quartile confidence errors (N=6, c0 >= 0.578):\nMCB: 0/6 (0.0%) | Entropy: 2/6 (33.3%)", "NOT MET (UNVERIFIED)", "No errors exist with c0 >= 0.90. On top-quartile confidence errors, MCB intercepted 0 of 6, while entropy caught 2 of 6. Hypothesis failed."],
        ["D4", "Direct fault-attribution accuracy exceeds random chance (>33.3%)", "Overall Accuracy: 17/60 = 28.3% (95% CI [18.5, 40.8%])\nBinomial p-value vs 1/3 chance: p = 0.25\nGeometric Collapse: 42 of 60 (70.0%)", "NOT MET", "Attribution accuracy is statistically indistinguishable from random chance. The decision rule collapsed to predicting Geometric 70% of the time."]
    ]
    
    t_rules = doc.add_table(rows=len(rules_table), cols=5)
    t_rules.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_vals in enumerate(rules_table):
        for c_idx, val in enumerate(row_vals):
            cell = t_rules.cell(r_idx, c_idx)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "FDECEC") # All failed / red
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # ── Section 4: Deep Dive into Attribution Failure & Grid Boundary ────────
    h4 = doc.add_heading("4. Analysis of Attribution Collapse & Boundary Artifacts", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    
    p_an = doc.add_paragraph()
    p_an.add_run("A critical examination of the failure modes reveals two profound methodological insights:\n\n")
    p_an.add_run("1. Why Attribution Collapsed (28.3% Accuracy): ").font.bold = True
    p_an.add_run("Across 60 synthetic corruptions, the argmax prediction collapsed to Geometric in 70% of cases. Rather than proving an interpretable failure mode, this reflects a biased decision boundary. Geometric transformations involve bilinear interpolation during rotation and zoom, which exacerbates pixel artifacts in already-degraded images. Furthermore, on correct images, Geometric has a slightly lower standard deviation (0.0748 vs 0.0841 for Sensor/Noise), artificially elevating z-scores. Semantic fault attribution was fundamentally unproven.\n\n")
    p_an.add_run("2. Boundary Artifact in Calibration: ").font.bold = True
    p_an.add_run("The calibrated tau_fam = 0.10 was located on the lowest boundary of the initial search grid. Extending the grid to include [0.025, 0.05, 0.075] revealed that Youden J peaks at tau_fam = 0.025 on Partition A. In this regime, the Cross-Family Spread (CFS) score trips on virtually any microscopic family noise, completely abandoning the multi-family compromise principle proposed in the original theory.")
    
    out_docx = Path("outputs/evaluation_report/Metamorphic_Circuit_Breaker_Results_Report_v4.docx")
    doc.save(str(out_docx))
    print(f"Saved unvarnished DOCX report to {out_docx}")

if __name__ == "__main__":
    create_honest_report()
